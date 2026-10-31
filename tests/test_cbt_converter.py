import io
import json
import pytest
import docx
from flask import Flask

from routes import acca_plus as plus


@pytest.fixture
def client():
    app = Flask(__name__, template_folder="../templates")
    app.config["SECRET_KEY"] = "test-secret"
    app.register_blueprint(plus.bp, url_prefix="/acca-plus")
    return app.test_client()


def test_base_template_exists():
    path = plus._get_base_cbt_template_path()
    assert path is not None
    doc = docx.Document(path)
    assert len(doc.tables) == 4


def test_parse_pg_and_pgk_questions():
    sample_text = """
    1. Siapakah presiden pertama Indonesia?
    A. Soekarno
    B. Soeharto
    C. Habibie
    D. Megawati
    Kunci: A
    Poin: 1

    2. Manakah yang termasuk planet dalam tata surya? (Pilih lebih dari satu)
    A. Mars
    B. Matahari
    C. Bumi
    D. Bulan
    Kunci: A, C
    Poin: 2
    """
    questions, summary = plus._parse_raw_questions(sample_text)
    assert len(questions) == 2
    assert summary["total"] == 2
    assert summary["missing_keys"] == 0

    # Q1
    assert questions[0]["no"] == 1
    assert questions[0]["type"] == "PG"
    assert questions[0]["key"] == "A"
    assert questions[0]["point"] == 1
    assert len(questions[0]["options"]) == 4

    # Q2
    assert questions[1]["no"] == 2
    assert questions[1]["type"] == "PGK"
    assert "A" in questions[1]["key"] and "C" in questions[1]["key"]
    assert questions[1]["point"] == 2


def test_parse_bs_and_matching_questions():
    sample_text = """
    1. Tentukan apakah pernyataan berikut Benar atau Salah:
    - Air mendidih pada suhu 100 C [B]
    - Es batu terasa panas saat disentuh [S]

    2. Pasangkan nama negara dengan ibu kotanya:
    - Indonesia -> Jakarta
    - Jepang -> Tokyo
    - Perancis -> Paris
    """
    questions, summary = plus._parse_raw_questions(sample_text)
    assert len(questions) == 2

    # Q1: BS
    assert questions[0]["type"] == "BS"
    assert len(questions[0]["statements"]) == 2
    assert questions[0]["statements"][0]["key"] == "B"
    assert questions[0]["statements"][1]["key"] == "S"
    assert questions[0]["key"] == "B|S"

    # Q2: MJ
    assert questions[1]["type"] == "MJ"
    assert len(questions[1]["pairs"]) == 3
    assert questions[1]["pairs"][0]["left"] == "Indonesia"
    assert questions[1]["pairs"][0]["right"] == "Jakarta"


def test_parse_isian_esai_and_ordering():
    sample_text = """
    1. Ibu kota negara Indonesia adalah .....
    Kunci: Jakarta

    2. Urutkan tahapan metode ilmiah berikut dari awal:
    1. Merumuskan masalah
    2. Membuat hipotesis
    3. Eksperimen
    4. Kesimpulan
    Kunci: 1;2;3;4

    3. Jelaskan proses terjadinya daur air secara singkat!
    """
    questions, summary = plus._parse_raw_questions(sample_text)
    assert len(questions) == 3

    # Q1: ISIAN
    assert questions[0]["type"] == "ISIAN"
    assert questions[0]["key"] == "Jakarta"

    # Q2: URUT
    assert questions[1]["type"] == "URUT"
    assert len(questions[1]["ordering_items"]) == 4
    assert questions[1]["key"] == "1;2;3;4"

    # Q3: ESAI
    assert questions[2]["type"] == "ESAI"
    assert questions[2]["point"] == 5


def test_parse_bottom_answer_keys():
    sample_text = """
    1. Siapakah presiden pertama RI?
    A. Soekarno
    B. Soeharto

    2. Ibu kota Perancis adalah .....

    KUNCI JAWABAN:
    1. A
    2. Paris
    """
    questions, summary = plus._parse_raw_questions(sample_text)
    assert len(questions) == 2
    assert questions[0]["key"] == "A"
    assert questions[1]["key"] == "Paris"
    assert summary["missing_keys"] == 0


def test_generate_cbt_docx():
    sample_questions = [
        {
            "no": 1,
            "type": "PG",
            "text": "Pertanyaan 1?",
            "key": "A",
            "point": 1,
            "options": [{"label": "A", "text": "Jawaban A"}, {"label": "B", "text": "Jawaban B"}],
            "kisi": {"materi": "Materi 1", "indikator": "Indikator 1"},
            "stimulus": {"title": "Stimulus 1", "content": "Bacaan teks 1..."},
        },
        {
            "no": 2,
            "type": "BS",
            "text": "Pernyataan benar salah?",
            "key": "B|S",
            "point": 2,
            "statements": [
                {"label": "P1", "text": "Pernyataan 1", "key": "B"},
                {"label": "P2", "text": "Pernyataan 2", "key": "S"},
            ],
        },
        {
            "no": 3,
            "type": "MJ",
            "text": "Jodohkan item berikut!",
            "key": "",
            "point": 2,
            "pairs": [{"left": "Kiri 1", "right": "Kanan 1"}, {"left": "Kiri 2", "right": "Kanan 2"}],
        },
    ]

    docx_stream = plus._generate_cbt_docx(sample_questions)
    assert docx_stream is not None
    assert len(docx_stream.getvalue()) > 0

    # Read back generated document
    doc = docx.Document(docx_stream)
    assert len(doc.tables) == 4

    # Table 0 (Soal)
    t0 = doc.tables[0]
    headers = [c.text.strip() for c in t0.rows[0].cells]
    assert headers[0] == "No"
    assert headers[1] == "Tipe"
    assert headers[2].startswith("Soal")
    assert headers[3] == "Kunci"
    assert headers[4] == "Poin"

    # Check question 1 row
    assert t0.rows[1].cells[0].text.strip() == "1"
    assert t0.rows[1].cells[1].text.strip() == "PG"
    assert t0.rows[1].cells[3].text.strip() == "A"

    # Table 1 (Kisi-kisi)
    t1 = doc.tables[1]
    assert len(t1.rows) >= 2

    # Table 2 (Stimulus)
    t2 = doc.tables[2]
    assert len(t2.rows) >= 2

    # Table 3 (Petunjuk)
    t3 = doc.tables[3]
    assert len(t3.rows) >= 2


def test_cbt_converter_routes(client):
    # GET page
    res_get = client.get("/acca-plus/cbt-converter")
    assert res_get.status_code == 200
    assert b"Konverter Soal" in res_get.data

    # POST action=download_template
    res_dl = client.post("/acca-plus/cbt-converter", data={"action": "download_template"})
    assert res_dl.status_code == 200
    assert "wordprocessingml" in res_dl.headers.get("Content-Type", "")

    # POST action=parse with raw_text
    sample_text = "1. Contoh soal?\nA. Opsi 1\nB. Opsi 2\nKunci: A\nPoin: 1"
    res_parse = client.post(
        "/acca-plus/cbt-converter",
        data={"action": "parse", "raw_text": sample_text},
    )
    assert res_parse.status_code == 200
    data = res_parse.get_json()
    assert data["success"] is True
    assert len(data["questions"]) == 1
    assert data["questions"][0]["key"] == "A"

    # POST action=generate with questions_json
    questions_json = json.dumps(data["questions"])
    res_gen = client.post(
        "/acca-plus/cbt-converter",
        data={"action": "generate", "questions_json": questions_json},
    )
    assert res_gen.status_code == 200
    assert "wordprocessingml" in res_gen.headers.get("Content-Type", "")
    assert len(res_gen.data) > 0
