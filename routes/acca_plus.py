import base64
import csv
import importlib.util
import io
import json
import os
import re
import zipfile

from flask import Blueprint, jsonify, render_template, request, send_file, url_for

from routes._helpers import NO_FILE_SINGLE, NO_FILE_MULTIPLE, log_error, safe_int, safe_float
from utils.pymupdf import import_pymupdf

bp = Blueprint("acca_plus", __name__)

fitz = import_pymupdf()

# ---------------------------------------------------------------------------
# Lazy pix2tex detection – heavy import is deferred until first request so
# the app starts instantly even when pix2tex is not installed.
# ---------------------------------------------------------------------------

HAS_PIX2TEX = importlib.util.find_spec("pix2tex") is not None

_latex_ocr_model = None  # cached singleton; loaded on first POST




def _get_model():
    """Return a cached LatexOCR instance, loading it on first call."""
    global _latex_ocr_model
    if _latex_ocr_model is None:
        from pix2tex.cli import LatexOCR  # noqa: PLC0415

        _latex_ocr_model = LatexOCR()
    return _latex_ocr_model


def _unique_name(name: str, used_names: set[str]) -> str:
    """Ensure filename is unique within a zip archive to prevent overwrites."""
    if name not in used_names:
        used_names.add(name)
        return name
    stem, ext = os.path.splitext(name)
    counter = 1
    while f"{stem}_{counter}{ext}" in used_names:
        counter += 1
    unique = f"{stem}_{counter}{ext}"
    used_names.add(unique)
    return unique


# ---------------------------------------------------------------------------
# Formula to LaTeX
# ---------------------------------------------------------------------------

@bp.route("/latex-ocr")
def latex_ocr_page():
    if HAS_PIX2TEX:
        status_html = (
            '<p><i class="bi bi-check-circle-fill" style="color:#16a34a"></i> '
            '<strong>pix2tex ready.</strong> LaTeX-OCR model detected.</p>'
        )
    else:
        status_html = (
            '<p><i class="bi bi-exclamation-triangle-fill" style="color:#ffb703"></i> '
            '<strong>pix2tex not installed.</strong> Run on the server: '
            '<code>pip install pix2tex</code> then restart. '
            'Requires PyTorch; model weights (~300 MB) are downloaded automatically on first use.</p>'
        )

    return render_template(
        "tools/latex_ocr.html",
        title="Formula to LaTeX",
        description="Convert a formula image to LaTeX code automatically",
        status_html=status_html,
        has_pix2tex=HAS_PIX2TEX,
        endpoint="/acca-plus/latex-ocr",
        accept=".jpg,.jpeg,.png,.bmp,.tiff,.webp",
    )


@bp.route("/latex-ocr", methods=["POST"])
def latex_ocr():
    if not HAS_PIX2TEX:
        return jsonify(
            error="pix2tex is not installed. Run: pip install pix2tex on the server, then restart."
        ), 400

    files = request.files.getlist("files")
    if not files or not files[0].filename:
        return jsonify(error=NO_FILE_SINGLE), 400

    from PIL import Image, ImageOps  # noqa: PLC0415

    try:
        raw_img = Image.open(io.BytesIO(files[0].read()))
        raw_img = ImageOps.exif_transpose(raw_img)
        if raw_img.mode in ("RGBA", "LA", "P"):
            bg = Image.new("RGB", raw_img.size, (255, 255, 255))
            rgba = raw_img.convert("RGBA")
            bg.paste(rgba, mask=rgba.split()[3])
            img = bg
        else:
            img = raw_img.convert("RGB")
    except Exception as e:
        log_error(e, "latex_ocr.open")
        return jsonify(error="Failed to open image. Make sure the file is a valid image."), 400

    try:
        model = _get_model()
        latex = model(img)
    except Exception as e:
        log_error(e, "latex_ocr.predict")
        return jsonify(error=f"LaTeX prediction failed: {e}"), 500

    if not latex or not latex.strip():
        return jsonify(latex="(No formula detected)")

    return jsonify(latex=latex.strip())








# ---------------------------------------------------------------------------
# Rename PDF by Content (NISN & Nama Siswa)
# ---------------------------------------------------------------------------

# ── Regex patterns ───────────────────────────────────────────────────────────
# NISN: exactly 10 digits, optionally preceded by a common label
_RE_NISN_LABELED = re.compile(
    r"NISN\s*[:\-]?\s*(\d{10})",
    re.IGNORECASE,
)
_RE_NISN_BARE = re.compile(r"\b(\d{10})\b")

# Nama Siswa: common label variants used in Indonesian school documents
_RE_NAMA = re.compile(
    r"(?:Nama\s+(?:Siswa|Peserta\s+Didik|Lengkap|Murid)|Nama)\s*[:\-]\s*([^\n\r,;]{3,80})",
    re.IGNORECASE,
)

# Guru / Pegawai (NIP & NUPTK):
_RE_NIP_LABELED = re.compile(r"NIP\s*[:\-]?\s*(\d{18})", re.IGNORECASE)
_RE_NIP_BARE = re.compile(r"\b((?:19|20)\d{16})\b")
_RE_NUPTK_LABELED = re.compile(r"NUPTK\s*[:\-]?\s*(\d{16})", re.IGNORECASE)
_RE_NUPTK_BARE = re.compile(r"\b(\d{16})\b")
_RE_NAMA_GURU = re.compile(
    r"(?:Nama\s+(?:Guru|Pegawai|PNS|GTK|Pendidik|Lengkap)|Nama)\s*[:\-]\s*([^\n\r,;]{3,80})",
    re.IGNORECASE,
)

# Surat Dinas / SK / Kwitansi:
_RE_NO_SURAT = re.compile(
    r"(?:Nomor\s+(?:Surat|SK|Naskah)|Nomor|No\.?\s*(?:Surat|SK)?)\s*[:\-]\s*([A-Za-z0-9\/\.\-_]{3,60})",
    re.IGNORECASE,
)
_RE_PERIHAL = re.compile(
    r"(?:Perihal|Hal|Tentang|Mata\s+Pelajaran|Mapel|Judul)\s*[:\-]\s*([^\n\r,;]{3,80})",
    re.IGNORECASE,
)


def _sanitize(text: str) -> str:
    """Remove characters that are invalid in filenames and trim whitespace."""
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text)
    text = re.sub(r"[\s_]+", "_", text.strip())
    return text[:80]


def _clean_nama(text: str) -> str:
    """Clean extracted student name, removing trailing fields, prefixes, and invalid chars."""
    if not text:
        return ""
    text = re.split(
        r"\b(?:NISN|NIS|Kelas|Tempat|TTL|Jenis Kelamin|L/P|Tahun|Semester|Agama|Kurikulum)\b",
        text,
        flags=re.IGNORECASE,
    )[0]
    text = re.sub(r"^[\s:\-\._#]+", "", text)
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:80]


def _extract_nisn_and_nama(text: str) -> tuple[str, str]:
    """Return (nisn, nama) extracted from PDF text, or empty strings."""
    nisn = ""
    m = _RE_NISN_LABELED.search(text)
    if m:
        nisn = m.group(1)
    else:
        m = _RE_NISN_BARE.search(text)
        if m:
            nisn = m.group(1)

    nama = ""
    m = _RE_NAMA.search(text)
    if m:
        nama = _sanitize(_clean_nama(m.group(1)))

    return nisn, nama


@bp.route("/rename-pdf")
def rename_pdf_page():
    return render_template(
        "tools/rename_pdf.html",
        title="Rename PDFs",
        description=(
            "Upload multiple PDFs at once. NISN and student name are extracted "
            "automatically from each document and used to rename the files: "
            "<code>filename_NISN_StudentName.pdf</code>"
        ),
        endpoint="/acca-plus/rename-pdf",
    )


@bp.route("/rename-pdf", methods=["POST"])
def rename_pdf():
    if fitz is None:
        return jsonify(error="PyMuPDF (fitz) is not installed on the server."), 500

    files = request.files.getlist("files")
    if not files or not files[0].filename:
        return jsonify(error=NO_FILE_MULTIPLE), 400

    custom_nisn_pattern = (request.form.get("nisn_pattern") or "").strip()
    custom_nama_pattern = (request.form.get("nama_pattern") or "").strip()

    try:
        re_nisn_custom = re.compile(custom_nisn_pattern, re.IGNORECASE) if custom_nisn_pattern else None
        re_nama_custom = re.compile(custom_nama_pattern, re.IGNORECASE) if custom_nama_pattern else None
    except re.error as e:
        return jsonify(error=f"Invalid regex pattern: {e}"), 400

    zip_buf = io.BytesIO()
    log_lines: list[str] = []
    used_names: set[str] = set()
    processed = 0

    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for upload in files:
            original_filename = upload.filename or "unknown.pdf"
            stem = os.path.splitext(original_filename)[0]
            data = b""

            try:
                data = upload.read()
                if not data:
                    raise ValueError("Uploaded file is empty.")
                doc = fitz.open(stream=data, filetype="pdf")

                pages_to_scan = min(3, doc.page_count)
                text = "\n".join(doc[i].get_text() for i in range(pages_to_scan))
                if not text.strip() and doc.page_count > 3:
                    text = "\n".join(page.get_text() for page in doc)
                doc.close()

                nisn = ""
                if re_nisn_custom:
                    m = re_nisn_custom.search(text)
                    nisn = m.group(1) if (m and m.lastindex) else (m.group(0) if m else "")
                else:
                    nisn, _ = _extract_nisn_and_nama(text)

                nama = ""
                if re_nama_custom:
                    m = re_nama_custom.search(text)
                    nama_raw = m.group(1) if (m and m.lastindex) else (m.group(0) if m else "")
                    nama = _sanitize(nama_raw)
                else:
                    _, nama = _extract_nisn_and_nama(text)

                parts = [stem]
                if nisn:
                    parts.append(nisn)
                if nama:
                    parts.append(nama)

                new_name = "_".join(parts) + ".pdf"
                status = "OK"
                if not nisn:
                    status += " [NISN not found]"
                if not nama:
                    status += " [Name not found]"

            except Exception as e:
                log_error(e, f"rename_pdf.{original_filename}")
                new_name = original_filename
                status = f"ERROR: {e}"

            unique_new_name = _unique_name(new_name, used_names)
            zf.writestr(unique_new_name, data)
            log_lines.append(f"{original_filename}  →  {unique_new_name}  [{status}]")
            processed += 1

        zf.writestr("_rename_log.txt", "\n".join(log_lines).encode("utf-8"))

    if processed == 0:
        return jsonify(error="No PDF files could be processed."), 400

    zip_buf.seek(0)
    return send_file(
        zip_buf,
        mimetype="application/zip",
        as_attachment=True,
        download_name="renamed_pdfs.zip",
    )


# ---------------------------------------------------------------------------
# Split PDF by Profile (Siswa, Guru/GTK, Surat Dinas, atau Label Bebas)
# ---------------------------------------------------------------------------

PROFILE_METADATA = {
    "siswa": {
        "label": "Siswa / Santri",
        "id_col": "NISN",
        "name_col": "Nama Siswa",
        "fallback_prefix": "Siswa",
    },
    "guru": {
        "label": "Guru / GTK / Pegawai",
        "id_col": "NIP / NUPTK",
        "name_col": "Nama Guru / Pegawai",
        "fallback_prefix": "Pegawai",
    },
    "surat": {
        "label": "Surat Dinas / SK / Kwitansi",
        "id_col": "Nomor Surat / SK",
        "name_col": "Perihal / Tentang",
        "fallback_prefix": "Surat",
    },
    "custom": {
        "label": "Kustom / Label Bebas",
        "id_col": "ID",
        "name_col": "Nama / Keterangan",
        "fallback_prefix": "Dokumen",
    },
}


def _extract_id_and_nama_by_profile(
    text: str,
    profile: str = "siswa",
    custom_label_id: str = "",
    custom_label_name: str = "",
    re_id_custom=None,
    re_name_custom=None,
) -> tuple[str, str]:
    """Extract (id_val, name_val) according to selected profile or custom labels."""
    id_val = ""
    name_val = ""

    # 1. Custom regex if provided has top priority
    if re_id_custom:
        m = re_id_custom.search(text)
        id_val = m.group(1) if (m and m.lastindex) else (m.group(0) if m else "")
    if re_name_custom:
        m = re_name_custom.search(text)
        raw_name = m.group(1) if (m and m.lastindex) else (m.group(0) if m else "")
        name_val = _clean_nama(raw_name)

    if id_val and name_val:
        return id_val.strip(), name_val.strip()

    # 2. Extract according to profile
    if profile == "guru":
        if not id_val:
            m = _RE_NIP_LABELED.search(text)
            if m:
                id_val = m.group(1)
            else:
                m = _RE_NIP_BARE.search(text)
                if m:
                    id_val = m.group(1)
                else:
                    m = _RE_NUPTK_LABELED.search(text) or _RE_NUPTK_BARE.search(text)
                    if m:
                        id_val = m.group(1)
        if not name_val:
            m = _RE_NAMA_GURU.search(text)
            if m:
                name_val = _clean_nama(m.group(1))

    elif profile == "surat":
        if not id_val:
            m = _RE_NO_SURAT.search(text)
            if m:
                id_val = m.group(1).strip()
        if not name_val:
            m = _RE_PERIHAL.search(text)
            if m:
                name_val = _clean_nama(m.group(1))

    elif profile == "custom":
        if custom_label_id and not id_val:
            pat = rf"(?:{re.escape(custom_label_id.strip())})\s*[:\-\. ]\s*([^\n\r,;]{{2,80}})"
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                id_val = m.group(1).strip()
        if custom_label_name and not name_val:
            pat = rf"(?:{re.escape(custom_label_name.strip())})\s*[:\-\. ]\s*([^\n\r,;]{{2,80}})"
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                name_val = _clean_nama(m.group(1))

    else:  # default "siswa"
        if not id_val:
            m = _RE_NISN_LABELED.search(text)
            if m:
                id_val = m.group(1)
            else:
                m = _RE_NISN_BARE.search(text)
                if m:
                    id_val = m.group(1)
        if not name_val:
            m = _RE_NAMA.search(text)
            if m:
                name_val = _clean_nama(m.group(1))

    return id_val.strip(), name_val.strip()


def _format_student_filename(
    index: int,
    start_page: int,
    end_page: int,
    nisn: str,
    nama: str,
    naming_pattern: str,
    prefix: str,
    uppercase_name: bool,
    used_names: set[str],
    profile: str = "siswa",
) -> str:
    """Format a clean, readable filename for a split PDF segment."""
    clean_id = _sanitize(nisn).strip("_")
    clean_name = _sanitize(nama).strip("_")
    if uppercase_name and clean_name:
        clean_name = clean_name.upper()

    clean_prefix = _sanitize(prefix).strip("_")
    meta = PROFILE_METADATA.get(profile, PROFILE_METADATA["siswa"])
    fallback_base = clean_prefix or meta["fallback_prefix"]
    fallback = f"{fallback_base}_{index:02d}_Hal_{start_page + 1}-{end_page}"

    if naming_pattern == "no_nisn_nama":
        parts = [f"{index:02d}"]
        if clean_prefix:
            parts.insert(0, clean_prefix)
        if clean_id:
            parts.append(clean_id)
        if clean_name:
            parts.append(clean_name)
        if not clean_id and not clean_name:
            parts.append(fallback)
        base = "_".join(parts)

    elif naming_pattern == "prefix_nisn_nama":
        pref = clean_prefix or fallback_base
        parts = [pref]
        if clean_id:
            parts.append(clean_id)
        if clean_name:
            parts.append(clean_name)
        if not clean_id and not clean_name:
            parts.append(fallback)
        base = "_".join(parts)

    elif naming_pattern == "nama_nisn":
        parts = []
        if clean_prefix:
            parts.append(clean_prefix)
        if clean_name:
            parts.append(clean_name)
        if clean_id:
            parts.append(clean_id)
        if not clean_id and not clean_name:
            parts.append(fallback)
        base = "_".join(parts)

    else:  # default: "nisn_nama" (ID_Name)
        parts = []
        if clean_prefix:
            parts.append(clean_prefix)
        if clean_id:
            parts.append(clean_id)
        if clean_name:
            parts.append(clean_name)
        if not clean_id and not clean_name:
            parts.append(fallback)
        base = "_".join(parts)

    filename = f"{base}.pdf"
    return _unique_name(filename, used_names)


def _detect_student_segments(
    doc,
    mode: str = "auto",
    profile: str = "siswa",
    custom_label_id: str = "",
    custom_label_name: str = "",
    fixed_pages: int = 1,
    keyword: str = "",
    re_nisn_custom=None,
    re_nama_custom=None,
    naming_pattern: str = "nisn_nama",
    prefix: str = "",
    uppercase_name: bool = True,
) -> tuple[list[dict], list[str]]:
    """Detect page ranges and identifiers in a bulk PDF based on selected profile.

    Returns: (segments_list, warnings_list)
    """
    total_pages = len(doc)
    if total_pages == 0:
        return [], ["Dokumen PDF kosong (0 halaman)."]

    warnings: list[str] = []
    meta = PROFILE_METADATA.get(profile, PROFILE_METADATA["siswa"])
    id_col = custom_label_id.strip() or meta["id_col"]
    name_col = custom_label_name.strip() or meta["name_col"]

    pages_meta = []
    has_any_text = False

    for i in range(total_pages):
        page = doc[i]
        text = page.get_text() or ""
        if text.strip():
            has_any_text = True

        id_val, name_val = _extract_id_and_nama_by_profile(
            text=text,
            profile=profile,
            custom_label_id=custom_label_id,
            custom_label_name=custom_label_name,
            re_id_custom=re_nisn_custom,
            re_name_custom=re_nama_custom,
        )

        pages_meta.append({
            "page_idx": i,
            "text": text,
            "id_val": id_val.strip(),
            "name_val": name_val.strip(),
        })

    if not has_any_text:
        warnings.append(
            "Tidak ditemukan teks digital pada PDF ini (kemungkinan dokumen hasil scan / gambar murni). "
            "Gunakan mode 'Jumlah Halaman Tetap' untuk memotong per halaman, atau jalankan OCR PDF terlebih dahulu."
        )

    # Determine boundary cut points (0-indexed page start indices)
    boundaries = [0]

    if mode == "fixed":
        fp = max(1, int(fixed_pages))
        for p in range(fp, total_pages, fp):
            boundaries.append(p)

    elif mode == "keyword":
        kw = (keyword or "").strip().lower()
        if not kw:
            warnings.append("Kata kunci pemisah kosong. Beralih ke deteksi otomatis.")
            mode = "auto"
        else:
            for i in range(1, total_pages):
                if kw in pages_meta[i]["text"].lower():
                    boundaries.append(i)

    if mode == "auto":
        cur_id = pages_meta[0]["id_val"]
        cur_name = pages_meta[0]["name_val"].lower()

        for i in range(1, total_pages):
            p_id = pages_meta[i]["id_val"]
            p_name = pages_meta[i]["name_val"].lower()

            is_new = False
            # Check 1: ID detected and changed
            if p_id and cur_id and p_id != cur_id:
                is_new = True
            # Check 2: ID appears after previous page didn't have one, with different name
            elif p_id and not cur_id and (p_name and p_name != cur_name):
                is_new = True
            # Check 3: Name changed
            elif p_name and cur_name and p_name != cur_name:
                is_new = True
            # Check 4: A new ID appears on page i after >= 1 page
            elif p_id and not cur_id and (i - boundaries[-1] >= 1):
                if p_name:
                    is_new = True

            if is_new:
                boundaries.append(i)
                cur_id = p_id
                cur_name = p_name
            else:
                if not cur_id and p_id:
                    cur_id = p_id
                if not cur_name and p_name:
                    cur_name = p_name

    boundaries.append(total_pages)
    clean_boundaries = []
    for b in boundaries:
        if not clean_boundaries or b > clean_boundaries[-1]:
            clean_boundaries.append(b)

    used_names: set[str] = set()
    segments = []

    for idx in range(len(clean_boundaries) - 1):
        start = clean_boundaries[idx]
        end = clean_boundaries[idx + 1]
        pages_count = end - start

        # Best ID and Name in this page span
        seg_id = ""
        seg_name = ""
        for p in range(start, end):
            if not seg_id and pages_meta[p]["id_val"]:
                seg_id = pages_meta[p]["id_val"]
            if not seg_name and pages_meta[p]["name_val"]:
                seg_name = pages_meta[p]["name_val"]
            if seg_id and seg_name:
                break

        out_name = _format_student_filename(
            index=idx + 1,
            start_page=start,
            end_page=end,
            nisn=seg_id,
            nama=seg_name,
            naming_pattern=naming_pattern,
            prefix=prefix,
            uppercase_name=uppercase_name,
            used_names=used_names,
            profile=profile,
        )

        display_range = f"Hal {start + 1}" if pages_count == 1 else f"Hal {start + 1} - {end}"

        status_notes = []
        if not seg_id:
            status_notes.append(f"{id_col} tdk ditemukan")
        if not seg_name:
            status_notes.append(f"{name_col} tdk ditemukan")
        status_text = ", ".join(status_notes) if status_notes else "Lengkap"

        formatted_name = seg_name.upper() if (uppercase_name and seg_name) else seg_name

        segments.append({
            "index": idx + 1,
            "start_page": start,
            "end_page": end,
            "pages_count": pages_count,
            "page_range_display": display_range,
            "nisn": seg_id,        # backward compatible alias
            "nama": formatted_name, # backward compatible alias
            "id_val": seg_id,
            "name_val": formatted_name,
            "filename": out_name,
            "status": status_text,
        })

    if len(segments) <= 1 and total_pages > 2 and mode == "auto":
        warnings.append(
            f"Hanya 1 entitas terdeteksi pada seluruh dokumen untuk profil '{meta['label']}'. "
            f"Pastikan dokumen memuat teks {id_col}/{name_col} atau gunakan mode 'Jumlah Halaman Tetap'."
        )

    return segments, warnings


def _split_student_pdf_to_zip(
    pdf_bytes: bytes,
    original_filename: str,
    segments: list[dict],
    mode_label: str,
    profile_label: str = "Siswa / Santri",
    id_col_label: str = "NISN",
    name_col_label: str = "Nama Siswa",
) -> tuple[bytes, str]:
    """Split the PDF into individual PDFs and pack into a ZIP archive with summary."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    zip_buf = io.BytesIO()

    stem = os.path.splitext(original_filename)[0]
    zip_filename = f"{_sanitize(stem)}_split_dokumen.zip"

    summary_lines = [
        "=" * 88,
        "                    ACCA TOOLS — BULK PDF SPLIT SUMMARY",
        "=" * 88,
        f"Dokumen Asal   : {original_filename}",
        f"Profil Dokumen : {profile_label}",
        f"Total Halaman  : {len(doc)}",
        f"Total File     : {len(segments)} file dihasilkan",
        f"Metode Split   : {mode_label}",
        "-" * 88,
        f"{'No.':<4} | {'Rentang Hal':<14} | {id_col_label:<14} | {name_col_label:<24} | {'Nama File Output'}",
        "-" * 88,
    ]

    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for seg in segments:
            start = seg["start_page"]
            end = seg["end_page"]
            filename = seg["filename"]

            subdoc = fitz.open()
            subdoc.insert_pdf(doc, from_page=start, to_page=end - 1)
            sub_bytes = subdoc.tobytes(garbage=3, deflate=True)
            subdoc.close()

            zf.writestr(filename, sub_bytes)

            id_disp = (seg.get("id_val") or seg.get("nisn") or "-")[:14]
            name_disp = (seg.get("name_val") or seg.get("nama") or "-")[:24]
            line = f"{seg['index']:<4} | {seg['page_range_display']:<14} | {id_disp:<14} | {name_disp:<24} | {filename}"
            summary_lines.append(line)

        summary_lines.append("=" * 88)
        zf.writestr("_split_summary.txt", "\n".join(summary_lines).encode("utf-8"))

    doc.close()
    zip_buf.seek(0)
    return zip_buf.getvalue(), zip_filename


@bp.route("/split-pdf")
def split_pdf_page():
    return render_template(
        "tools/split_student_pdf.html",
        title="Split Document",
        description="Pecah dokumen PDF gabungan secara otomatis berdasarkan profil Siswa, Guru, atau Label Kustom.",
        endpoint="/acca-plus/split-pdf",
        accept=".pdf",
    )


@bp.route("/split-pdf", methods=["POST"])
def split_pdf():
    if fitz is None:
        return jsonify(error="PyMuPDF (fitz) is not installed on the server."), 500

    files = request.files.getlist("files")
    if not files or not files[0].filename:
        return jsonify(error=NO_FILE_SINGLE), 400

    action = (request.form.get("action") or "split").strip().lower()
    doc_profile = (request.form.get("doc_profile") or "siswa").strip().lower()
    custom_label_id = (request.form.get("custom_label_id") or "").strip()
    custom_label_name = (request.form.get("custom_label_name") or "").strip()

    mode = (request.form.get("mode") or "auto").strip().lower()
    fixed_pages = safe_int(request.form.get("fixed_pages"), default=1, min_val=1, max_val=500)
    keyword = (request.form.get("keyword") or "").strip()
    naming_pattern = (request.form.get("naming_pattern") or "nisn_nama").strip().lower()
    prefix = (request.form.get("prefix") or "").strip()
    if "uppercase_name" in request.form:
        uppercase_name = request.form.get("uppercase_name") in ("1", "true", "on", True)
    else:
        uppercase_name = True

    custom_nisn_pattern = (request.form.get("nisn_pattern") or "").strip()
    custom_nama_pattern = (request.form.get("nama_pattern") or "").strip()

    try:
        re_nisn_custom = re.compile(custom_nisn_pattern, re.IGNORECASE) if custom_nisn_pattern else None
        re_nama_custom = re.compile(custom_nama_pattern, re.IGNORECASE) if custom_nama_pattern else None
    except re.error as e:
        return jsonify(error=f"Pola regex tidak valid: {e}"), 400

    upload = files[0]
    filename = upload.filename or "bulk_documents.pdf"
    try:
        pdf_bytes = upload.read()
        if not pdf_bytes:
            return jsonify(error="File yang diunggah kosong."), 400

        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        if doc.is_encrypted:
            doc.close()
            return jsonify(error="File PDF terenkripsi/password. Harap buka kunci terlebih dahulu."), 400

        total_pages = len(doc)
        if total_pages == 0:
            doc.close()
            return jsonify(error="File PDF tidak memiliki halaman."), 400

        segments, warnings = _detect_student_segments(
            doc,
            mode=mode,
            profile=doc_profile,
            custom_label_id=custom_label_id,
            custom_label_name=custom_label_name,
            fixed_pages=fixed_pages,
            keyword=keyword,
            re_nisn_custom=re_nisn_custom,
            re_nama_custom=re_nama_custom,
            naming_pattern=naming_pattern,
            prefix=prefix,
            uppercase_name=uppercase_name,
        )
        doc.close()

    except Exception as e:
        log_error(e, f"split_pdf.parse.{filename}")
        return jsonify(error=f"Gagal memproses dokumen PDF: {e}"), 400

    if not segments:
        return jsonify(error="Tidak ada segmen dokumen yang dapat dipisahkan."), 400

    meta = PROFILE_METADATA.get(doc_profile, PROFILE_METADATA["siswa"])
    id_col_label = custom_label_id or meta["id_col"]
    name_col_label = custom_label_name or meta["name_col"]
    profile_label = meta["label"]

    if action == "preview":
        return jsonify(
            success=True,
            filename=filename,
            total_pages=total_pages,
            total_students=len(segments),
            segments=segments,
            warnings=warnings,
            mode=mode,
            doc_profile=doc_profile,
            id_col_label=id_col_label,
            name_col_label=name_col_label,
            profile_label=profile_label,
        )

    mode_labels = {
        "auto": f"Auto-detect Pergantian ({id_col_label} / {name_col_label})",
        "fixed": f"Jumlah Halaman Tetap ({fixed_pages} hal/file)",
        "keyword": f"Kata Kunci Halaman Pertama ('{keyword}')",
    }
    mode_label = mode_labels.get(mode, mode)

    try:
        zip_bytes, zip_filename = _split_student_pdf_to_zip(
            pdf_bytes=pdf_bytes,
            original_filename=filename,
            segments=segments,
            mode_label=mode_label,
            profile_label=profile_label,
            id_col_label=id_col_label,
            name_col_label=name_col_label,
        )
    except Exception as e:
        log_error(e, f"split_pdf.generate_zip.{filename}")
        return jsonify(error=f"Gagal membuat file ZIP: {e}"), 500

    return send_file(
        io.BytesIO(zip_bytes),
        mimetype="application/zip",
        as_attachment=True,
        download_name=zip_filename,
    )


# ---------------------------------------------------------------------------
# Bulk Certificate / Piagam Generator
# ---------------------------------------------------------------------------

def _hex_to_rgb(hex_str: str) -> tuple[float, float, float]:
    """Convert hex color string like #1a202c to RGB float tuple (0.0 to 1.0)."""
    s = (hex_str or "").strip().lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    if len(s) != 6:
        return (0.0, 0.0, 0.0)
    try:
        r = int(s[0:2], 16) / 255.0
        g = int(s[2:4], 16) / 255.0
        b = int(s[4:6], 16) / 255.0
        return (r, g, b)
    except ValueError:
        return (0.0, 0.0, 0.0)


def _render_text_template(template_str: str, row_data: dict) -> str:
    """Substitute {{Column}} and {Column} placeholders from row_data case-insensitively."""
    res = template_str or ""
    for col, val in row_data.items():
        val_str = str(val).strip() if val is not None else ""
        res = re.sub(rf"\{{\{{\s*{re.escape(col)}\s*\}}\}}", val_str, res, flags=re.IGNORECASE)
        res = re.sub(rf"\{{\s*{re.escape(col)}\s*\}}", val_str, res, flags=re.IGNORECASE)
    return res


def _parse_spreadsheet_rows(file_bytes: bytes, filename: str) -> tuple[list[str], list[dict]]:
    """Parse an Excel (.xlsx / .xls) or CSV file and return (headers, rows_as_dicts)."""
    ext = os.path.splitext(filename)[1].lower()
    headers: list[str] = []
    rows: list[dict] = []

    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            return [], []

        raw_headers = raw_rows[0]
        headers = [str(h).strip() for h in raw_headers if h is not None and str(h).strip()]

        for r in raw_rows[1:]:
            if not any(val is not None and str(val).strip() for val in r):
                continue
            row_dict = {}
            for col_idx, h in enumerate(headers):
                val = r[col_idx] if col_idx < len(r) else ""
                row_dict[h] = str(val).strip() if val is not None else ""
            rows.append(row_dict)

    else:  # CSV
        text_content = ""
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text_content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        reader = csv.reader(io.StringIO(text_content))
        raw_rows = list(reader)
        if not raw_rows:
            return [], []

        raw_headers = raw_rows[0]
        headers = [h.strip() for h in raw_headers if h.strip()]

        for r in raw_rows[1:]:
            if not any(val.strip() for val in r):
                continue
            row_dict = {}
            for col_idx, h in enumerate(headers):
                val = r[col_idx] if col_idx < len(r) else ""
                row_dict[h] = val.strip()
            rows.append(row_dict)

    return headers, rows


def _render_single_certificate(
    template_bytes: bytes,
    is_pdf_template: bool,
    row_data: dict,
    fields: list[dict],
    qr_config: dict,
) -> fitz.Document:
    """Generate a 1-page fitz.Document for a single recipient."""
    if is_pdf_template:
        tpl_doc = fitz.open(stream=template_bytes, filetype="pdf")
        doc = fitz.open()
        doc.insert_pdf(tpl_doc, from_page=0, to_page=0)
        tpl_doc.close()
        page = doc[0]
        page_w = page.rect.width
        page_h = page.rect.height
    else:
        from PIL import Image
        img = Image.open(io.BytesIO(template_bytes))
        page_w, page_h = img.width, img.height
        doc = fitz.open()
        page = doc.new_page(width=page_w, height=page_h)
        page.insert_image(page.rect, stream=template_bytes)

    # 1. Stamp text fields
    for field in fields:
        raw_text = field.get("text", "")
        rendered = _render_text_template(raw_text, row_data)
        if not rendered.strip():
            continue

        y_pct = safe_float(field.get("y"), default=50.0) / 100.0
        fontsize = safe_float(field.get("size"), default=24.0)
        fontname = field.get("font", "hebo")
        color = _hex_to_rgb(field.get("color", "#000000"))

        y_center = y_pct * page_h
        box_h = max(30.0, fontsize * 2.2)
        y0 = max(0.0, y_center - box_h / 2.0)
        y1 = min(page_h, y_center + box_h / 2.0)
        margin_x = page_w * 0.06
        rect = fitz.Rect(margin_x, y0, page_w - margin_x, y1)

        page.insert_textbox(
            rect,
            rendered,
            fontsize=fontsize,
            fontname=fontname,
            color=color,
            align=fitz.TEXT_ALIGN_CENTER,
        )

    # 2. Stamp QR code if enabled
    if qr_config.get("enabled"):
        qr_raw = qr_config.get("content", "")
        qr_content = _render_text_template(qr_raw, row_data)
        if qr_content.strip():
            import qrcode
            qr = qrcode.QRCode(
                version=None,
                error_correction=qrcode.constants.ERROR_CORRECT_M,
                box_size=8,
                border=1,
            )
            qr.add_data(qr_content)
            qr.make(fit=True)
            qr_img = qr.make_image(fill_color="black", back_color="white")
            qr_buf = io.BytesIO()
            qr_img.save(qr_buf, format="PNG")
            qr_bytes = qr_buf.getvalue()

            qr_size = safe_float(qr_config.get("size"), default=70.0)
            pos = qr_config.get("position", "bottom-left")

            if pos == "bottom-right":
                x0 = page_w * 0.94 - qr_size
                y0 = page_h * 0.92 - qr_size
            elif pos == "bottom-center":
                x0 = (page_w - qr_size) / 2.0
                y0 = page_h * 0.92 - qr_size
            else:  # bottom-left default
                x0 = page_w * 0.06
                y0 = page_h * 0.92 - qr_size

            qr_rect = fitz.Rect(x0, y0, x0 + qr_size, y0 + qr_size)
            page.insert_image(qr_rect, stream=qr_bytes)

    return doc


@bp.route("/certificate-generator")
def certificate_generator_page():
    return render_template(
        "tools/certificate_generator.html",
        title="Certificate Generator",
        description="Buat sertifikat atau piagam secara masal dari template dan data Excel/CSV.",
        endpoint="/acca-plus/certificate-generator",
    )


@bp.route("/certificate-generator", methods=["POST"])
def certificate_generator():
    if fitz is None:
        return jsonify(error="PyMuPDF (fitz) is not installed on the server."), 500

    action = (request.form.get("action") or "generate").strip().lower()

    # Parse Excel/CSV header scan only
    if action == "parse_data":
        data_file = request.files.get("data_file")
        if not data_file or not data_file.filename:
            return jsonify(error="Silakan unggah file data Excel atau CSV."), 400
        try:
            data_bytes = data_file.read()
            headers, rows = _parse_spreadsheet_rows(data_bytes, data_file.filename)
            if not headers:
                return jsonify(error="File Excel/CSV tidak memiliki kolom header."), 400
            return jsonify(
                success=True,
                headers=headers,
                total_rows=len(rows),
                sample_row=rows[0] if rows else {},
            )
        except Exception as e:
            log_error(e, "certificate_generator.parse_data")
            return jsonify(error=f"Gagal membaca data Excel/CSV: {e}"), 400

    # For preview and generate, we require template_file and data_file
    template_file = request.files.get("template_file")
    data_file = request.files.get("data_file")

    if not template_file or not template_file.filename:
        return jsonify(error="Silakan unggah file template sertifikat (.pdf, .png, .jpg)."), 400
    if not data_file or not data_file.filename:
        return jsonify(error="Silakan unggah file data penerima (.xlsx, .xls, .csv)."), 400

    template_filename = template_file.filename
    is_pdf_template = template_filename.lower().endswith(".pdf")

    try:
        template_bytes = template_file.read()
        if not template_bytes:
            return jsonify(error="File template kosong."), 400

        data_bytes = data_file.read()
        headers, rows = _parse_spreadsheet_rows(data_bytes, data_file.filename)
        if not rows:
            return jsonify(error="File data Excel/CSV tidak memuat baris data penerima."), 400

        # Parse fields config
        fields_json = request.form.get("fields_json") or "[]"
        try:
            fields = json.loads(fields_json)
        except Exception:
            fields = []

        if not fields:
            # Default fallback fields
            fields = [
                {"text": "{{Nama}}", "y": 42.0, "size": 32.0, "font": "hebo", "color": "#1a202c"},
                {"text": "Sebagai {{Sebagai}}", "y": 52.0, "size": 18.0, "font": "hebo", "color": "#333333"},
                {"text": "No: {{No_Piagam}}", "y": 24.0, "size": 13.0, "font": "hebo", "color": "#4a5568"},
            ]

        # Parse QR config
        qr_enabled = request.form.get("qr_enabled") in ("1", "true", "on", True)
        qr_content = (request.form.get("qr_content") or "").strip()
        qr_position = (request.form.get("qr_position") or "bottom-left").strip()
        qr_size = safe_float(request.form.get("qr_size"), default=70.0, min_val=30.0, max_val=200.0)

        qr_config = {
            "enabled": qr_enabled,
            "content": qr_content,
            "position": qr_position,
            "size": qr_size,
        }

    except Exception as e:
        log_error(e, "certificate_generator.init")
        return jsonify(error=f"Gagal menyiapkan template dan data: {e}"), 400

    # ── Action: Preview ────────────────────────────────────────────────────────
    if action == "preview":
        try:
            sample_row = rows[0]
            doc = _render_single_certificate(
                template_bytes=template_bytes,
                is_pdf_template=is_pdf_template,
                row_data=sample_row,
                fields=fields,
                qr_config=qr_config,
            )
            pix = doc[0].get_pixmap(dpi=115)
            png_bytes = pix.tobytes("png")
            doc.close()

            b64_img = base64.b64encode(png_bytes).decode("ascii")
            return jsonify(
                success=True,
                image=f"data:image/png;base64,{b64_img}",
                total_rows=len(rows),
                sample_row=sample_row,
                headers=headers,
            )
        except Exception as e:
            log_error(e, "certificate_generator.preview")
            return jsonify(error=f"Gagal membuat pratinjau: {e}"), 500

    # ── Action: Generate ───────────────────────────────────────────────────────
    output_format = (request.form.get("output_format") or "zip").strip().lower()

    try:
        combined_doc = fitz.open()
        zip_buf = io.BytesIO()
        used_names: set[str] = set()

        summary_lines = [
            "=" * 85,
            "               ACCA TOOLS — BULK CERTIFICATE GENERATION SUMMARY",
            "=" * 85,
            f"Template Asal   : {template_filename}",
            f"File Data       : {data_file.filename}",
            f"Total Penerima  : {len(rows)} orang",
            f"QR Verifikasi   : {'Aktif' if qr_enabled else 'Nonaktif'}",
            "-" * 85,
            f"{'No.':<4} | {'Nama Penerima':<30} | {'No. Piagam / ID':<20} | {'Nama File Output'}",
            "-" * 85,
        ]

        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for i, row in enumerate(rows):
                single_doc = _render_single_certificate(
                    template_bytes=template_bytes,
                    is_pdf_template=is_pdf_template,
                    row_data=row,
                    fields=fields,
                    qr_config=qr_config,
                )

                # Add to merged PDF
                combined_doc.insert_pdf(single_doc, from_page=0, to_page=0)

                # Determine file name
                name_val = row.get("Nama") or row.get("Nama Siswa") or row.get("Nama Guru") or f"Penerima_{i+1}"
                id_val = row.get("No_Piagam") or row.get("Nomor") or row.get("NISN") or ""

                clean_name = _sanitize(str(name_val))
                clean_id = _sanitize(str(id_val))

                parts = [f"{i+1:02d}"]
                if clean_id:
                    parts.append(clean_id)
                if clean_name:
                    parts.append(clean_name)

                out_filename = f"{'_'.join(parts)}.pdf"
                out_filename = _unique_name(out_filename, used_names)

                pdf_bytes = single_doc.tobytes(garbage=3, deflate=True)
                single_doc.close()

                zf.writestr(out_filename, pdf_bytes)

                line = f"{i+1:<4} | {str(name_val)[:30]:<30} | {str(id_val)[:20]:<20} | {out_filename}"
                summary_lines.append(line)

            summary_lines.append("=" * 85)

            # Insert combined PDF inside the zip too
            combined_bytes = combined_doc.tobytes(garbage=3, deflate=True)
            zf.writestr("_Semua_Piagam_Gabungan.pdf", combined_bytes)
            zf.writestr("_certificate_summary.txt", "\n".join(summary_lines).encode("utf-8"))

        combined_doc.close()

    except Exception as e:
        log_error(e, "certificate_generator.generate")
        return jsonify(error=f"Gagal men-generate sertifikat: {e}"), 500

    # If user explicitly requested single consolidated PDF only:
    if output_format == "single_pdf":
        return send_file(
            io.BytesIO(combined_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name="Semua_Piagam_Gabungan.pdf",
        )

    # Default: ZIP archive containing individual PDFs + merged PDF + summary
    zip_buf.seek(0)
    return send_file(
        zip_buf,
        mimetype="application/zip",
        as_attachment=True,
        download_name="piagam_sertifikat_masal.zip",
    )


# ---------------------------------------------------------------------------
# Compress PDF to Target Size
# ---------------------------------------------------------------------------

def _compress_pdf_at(data: bytes, jpeg_quality: int, max_dpi: int) -> bytes:
    """Optimize embedded images at their actual printed size, preserving masks.

    Text and vector paths are never rasterized. Shared logos are rewritten once
    at the largest size at which they appear anywhere in the document.
    """
    import math
    from PIL import Image

    with fitz.open(stream=data, filetype="pdf") as doc:
        images = {}
        masks = set()
        for page in doc:
            for info in page.get_images(full=True):
                xref, mask = info[:2]
                if mask:
                    masks.add(mask)
                item = images.setdefault(xref, {"page": page.number, "mask": mask,
                                                "width": 0.0, "height": 0.0})
                for rect in page.get_image_rects(xref):
                    item["width"] = max(item["width"], rect.width)
                    item["height"] = max(item["height"], rect.height)
        for xref, item in images.items():
            if xref in masks:
                continue
            try:
                # Color-key masks and stencils require separate handling.
                if doc.xref_get_key(xref, "Mask")[0] != "null":
                    continue
                if doc.xref_get_key(xref, "ImageMask")[1] == "true":
                    continue
                pix = fitz.Pixmap(doc, xref)
                if pix.colorspace is None:
                    continue
                if pix.colorspace.n >= 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                if item["mask"]:
                    pix = fitz.Pixmap(pix, fitz.Pixmap(doc, item["mask"]))
                with Image.open(io.BytesIO(pix.tobytes("png"))) as original:
                    image = original.copy()
                if max_dpi and item["width"] and item["height"]:
                    width = max(1, math.ceil(item["width"] * max_dpi / 72))
                    height = max(1, math.ceil(item["height"] * max_dpi / 72))
                    scale = min(1.0, max(width / image.width, height / image.height))
                    if scale < 1:
                        image = image.resize((max(1, math.ceil(image.width * scale)),
                                              max(1, math.ceil(image.height * scale))),
                                             Image.Resampling.LANCZOS)
                png = io.BytesIO()
                image.save(png, format="PNG", optimize=True)
                candidates = [png.getvalue()]
                # PNG retains alpha and avoids JPEG artifacts around logos/text.
                if image.mode not in ("RGBA", "LA", "P"):
                    jpeg = io.BytesIO()
                    image.save(jpeg, format="JPEG", quality=jpeg_quality,
                               optimize=False)
                    candidates.append(jpeg.getvalue())
                replacement = min(candidates, key=len)
                original_size = len(doc.xref_stream_raw(xref) or b"")
                if item["mask"]:
                    original_size += len(doc.xref_stream_raw(item["mask"]) or b"")
                if len(replacement) < original_size:
                    dummy = doc.new_page()
                    try:
                        new_xref = dummy.insert_image(dummy.rect, stream=replacement)
                        if new_xref != xref:
                            doc.xref_copy(new_xref, xref)
                    finally:
                        doc.delete_page(dummy.number)
            except Exception as exc:
                log_error(exc, f"compress_pdf image={xref}")
        return doc.tobytes(garbage=4, deflate=True,
                           deflate_images=True, deflate_fonts=True,
                           use_objstms=1, compression_effort=100)


def _compress_with_pikepdf(data: bytes) -> bytes:
    """Lossless compression using pikepdf object stream compression."""
    if importlib.util.find_spec("pikepdf") is None:
        return data

    try:
        import pikepdf
        with pikepdf.open(io.BytesIO(data)) as pdf:
            out = io.BytesIO()
            pdf.save(out, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate)
            return out.getvalue()
    except Exception as e:
        log_error(e, "_compress_with_pikepdf")
        return data



def _lossless_pdf(data: bytes) -> bytes:
    """Optimize objects/fonts without rasterizing text or re-encoding images."""
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            # Preserve font programs: re-subsetting existing subset fonts can
            # change space metrics and text selection in ReportLab documents.
            optimized = doc.tobytes(garbage=4, deflate=True, deflate_images=True,
                                   deflate_fonts=True, use_objstms=1,
                                   compression_effort=100)
    except Exception as e:
        log_error(e, "lossless_pdf")
        optimized = data
    candidates = [data, optimized, _compress_with_pikepdf(optimized)]
    return min(candidates, key=len)


def _compress_to_target(data: bytes, target_bytes: int, force_render: bool = False,
                        lossless: bool = False) -> tuple[bytes, int, bool]:
    """Find an output within the target while retaining text and vector tables.

    force_render is a legacy API flag, now enabling the same safe image-only
    optimization. Whole-page rasterization is deliberately never performed.
    """
    if target_bytes <= 0:
        raise ValueError("Target must be positive.")
    if len(data) <= target_bytes:
        return data, len(data), True
    best = _lossless_pdf(data)
    if lossless or len(best) <= target_bytes:
        return best, len(best), len(best) <= target_bytes
    base = best
    with fitz.open(stream=base, filetype="pdf") as doc:
        images = {info[0]: info for page in doc for info in page.get_images(full=True)}
    if not images:
        return best, len(best), False
    # Begin at print resolution: small logos often have thousands of pixels
    # although they occupy only a few centimetres on the page.
    has_opaque = any(not info[1] for info in images.values())
    for dpi in (600, 450, 300, 240, 200, 150):
        trial = _compress_pdf_at(base, 95, dpi)
        if len(trial) < len(best):
            best = trial
        if len(trial) <= target_bytes:
            return trial, len(trial), True
        if has_opaque:
            # Do not sacrifice readability by driving JPEG quality toward zero.
            lo, hi = 60, 94
            fitting = None
            while lo <= hi:
                quality = (lo + hi) // 2
                trial = _compress_pdf_at(base, quality, dpi)
                if len(trial) < len(best):
                    best = trial
                if len(trial) <= target_bytes:
                    fitting = trial
                    lo = quality + 1
                else:
                    hi = quality - 1
            if fitting is not None:
                return fitting, len(fitting), True
    return best, len(best), len(best) <= target_bytes


@bp.route("/compress-pdf")
def compress_pdf_page():
    return render_template(
        "tools/compress_pdf_plus.html",
        title="Compress PDF to Target Size",
        description=(
            "Upload a PDF and set your desired output file size. "
            "The tool automatically finds the best compression settings "
            "to get as close as possible to your target."
        ),
        endpoint="/acca-plus/compress-pdf",
    )


@bp.route("/compress-pdf", methods=["POST"])
def compress_pdf_target():
    if fitz is None:
        return jsonify(error="PyMuPDF (fitz) is not installed on the server."), 500

    files = request.files.getlist("files")
    if not files or not files[0].filename:
        return jsonify(error="Please upload a PDF file."), 400

    # Parse target size from form
    try:
        target_value = float(request.form.get("target_size", "0") or "0")
        target_unit = (request.form.get("target_unit", "KB") or "KB").upper()
        target_bytes = int(target_value * (1024 * 1024 if target_unit == "MB" else 1024))
    except (ValueError, TypeError, OverflowError):
        return jsonify(error="Invalid target size value."), 400

    if target_bytes <= 0:
        return jsonify(error="Target size must be greater than 0."), 400

    mode = request.form.get("mode") or ("render" if request.form.get("force_render") == "1" else "smart")
    if mode not in ("smart", "lossless", "balanced", "render"):
        return jsonify(error="Invalid compression mode."), 400
    if target_unit not in ("KB", "MB"):
        return jsonify(error="Invalid target size unit."), 400

    upload = files[0]
    stem = os.path.splitext(upload.filename or "document.pdf")[0]

    try:
        data = upload.read()
    except Exception as e:
        log_error(e, "compress_pdf_target.read")
        return jsonify(error="Failed to read uploaded file."), 400

    if not data:
        return jsonify(error="Uploaded file is empty."), 400

    try:
        chk = fitz.open(stream=data, filetype="pdf")
        if chk.needs_pass or chk.page_count == 0:
            chk.close()
            return jsonify(error="Please upload a non-empty, unlocked PDF."), 400
        chk.close()
    except Exception:
        return jsonify(error="Could not open PDF. File may be corrupted or encrypted."), 400

    original_size = len(data)

    if request.form.get("async") == "1":
        from utils.pdf_jobs import start_job
        try:
            token = start_job(data, target_bytes, mode, f"{stem}_compressed.pdf")
        except RuntimeError as e:
            return jsonify(error=str(e)), 503
        except Exception as e:
            log_error(e, "compress_pdf_target.start")
            return jsonify(error="Unable to start PDF compression. Please retry."), 500
        return jsonify(status_url=url_for("acca_plus.compress_pdf_status", token=token)), 202

    try:
        result_data, achieved_size, target_met = _compress_to_target(
            data, target_bytes, force_render=mode == "render", lossless=mode == "lossless")
    except Exception as e:
        log_error(e, "compress_pdf_target.process")
        return jsonify(error=f"Compression failed: {e}"), 500

    out_buf = io.BytesIO(result_data)
    out_buf.seek(0)

    resp = send_file(
        out_buf,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"{stem}_compressed.pdf",
    )
    # Expose size metadata via response headers for the JS result display
    resp.headers["X-Original-Size"] = str(original_size)
    resp.headers["X-Achieved-Size"] = str(achieved_size)
    resp.headers["X-Target-Size"] = str(target_bytes)
    resp.headers["X-Target-Met"] = "1" if target_met else "0"
    resp.headers["Access-Control-Expose-Headers"] = (
        "X-Original-Size, X-Achieved-Size, X-Target-Size, X-Target-Met"
    )
    return resp


@bp.route("/compress-pdf/jobs/<token>")
def compress_pdf_status(token):
    from utils.pdf_jobs import status
    result = status(token)
    if result is None:
        return jsonify(error="Compression result expired or was not found."), 404
    if result["state"] == "done":
        result["download_url"] = url_for("acca_plus.compress_pdf_download", token=token)
    response = jsonify(result)
    response.headers["Cache-Control"] = "no-store"
    return response


@bp.route("/compress-pdf/jobs/<token>/download")
def compress_pdf_download(token):
    from utils.pdf_jobs import job_path, status
    result = status(token)
    if result is None:
        return jsonify(error="Compression result expired or was not found."), 404
    if result["state"] != "done":
        return jsonify(error="Compression is not complete."), 409
    path = job_path(token)
    settings = json.loads((path / "settings.json").read_text(encoding="utf-8"))
    response = send_file(path / "output.pdf", mimetype="application/pdf",
                         as_attachment=True, download_name=settings["filename"])
    for name, value in {"Original-Size": settings["original"],
                        "Achieved-Size": result["achieved"],
                        "Target-Size": settings["target"],
                        "Target-Met": int(result["target_met"])}.items():
        response.headers[f"X-{name}"] = str(value)
    response.headers["Cache-Control"] = "no-store"
    return response


# ---------------------------------------------------------------------------
# Compress Image to Target Size
# ---------------------------------------------------------------------------

def _compress_image_to_target(img_data: bytes, target_bytes: int) -> tuple[bytes, int, bool, str, str]:
    """
    Binary-search compression settings to hit target_bytes for a single image.
    Preserves alpha transparency (using PNG or WebP) when the input has transparency,
    instead of replacing it with a solid white background.
    Opaque images are compressed to JPEG.
    Returns (best_data, achieved_size, target_met, extension, mimetype).
    """
    from PIL import Image as PILImage, ImageOps

    if not img_data:
        raise ValueError("Image file is empty.")
    if target_bytes <= 0:
        raise ValueError("Target size must be greater than 0.")

    try:
        pil_raw = PILImage.open(io.BytesIO(img_data))
        orig_fmt = (pil_raw.format or "JPEG").upper()
        pil_img = ImageOps.exif_transpose(pil_raw)
    except Exception as e:
        raise ValueError(f"Invalid image: {e}")

    def _fmt_info(fmt: str) -> tuple[str, str]:
        if fmt == "PNG":
            return ".png", "image/png"
        elif fmt == "WEBP":
            return ".webp", "image/webp"
        elif fmt == "BMP":
            return ".bmp", "image/bmp"
        return ".jpg", "image/jpeg"

    original_size = len(img_data)
    orig_ext, orig_mime = _fmt_info(orig_fmt)

    # Detect transparency (alpha channel)
    has_alpha = False
    if pil_img.mode in ("RGBA", "LA"):
        alpha = pil_img.getchannel("A")
        if alpha.getextrema()[0] < 255:
            has_alpha = True
        else:
            pil_img = pil_img.convert("RGB")
    elif pil_img.mode == "P":
        if "transparency" in pil_img.info:
            has_alpha = True
            pil_img = pil_img.convert("RGBA")
        else:
            pil_img = pil_img.convert("RGB")
    elif pil_img.mode not in ("RGB", "L"):
        pil_img = pil_img.convert("RGB")

    # If original file already meets target and didn't need rotation, preserve directly
    if original_size <= target_bytes and pil_raw.size == pil_img.size:
        return img_data, original_size, True, orig_ext, orig_mime

    orig_w, orig_h = pil_img.size

    def _render(fmt: str, quality: int, scale: float) -> bytes:
        img = pil_img
        if scale < 1.0:
            w = max(1, int(orig_w * scale))
            h = max(1, int(orig_h * scale))
            img = pil_img.resize((w, h), PILImage.LANCZOS)
        buf = io.BytesIO()
        if fmt == "JPEG":
            if img.mode in ("RGBA", "LA", "P"):
                bg = PILImage.new("RGB", img.size, (255, 255, 255))
                rgba = img.convert("RGBA")
                bg.paste(rgba, mask=rgba.split()[3])
                img = bg
            elif img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            img.save(buf, format="JPEG", quality=quality, optimize=True)
        elif fmt == "WEBP":
            img.save(buf, format="WEBP", quality=quality)
        elif fmt == "PNG":
            img.save(buf, format="PNG", optimize=True)
        return buf.getvalue()

    best_data = b""
    best_size = float("inf")
    best_ext = ".jpg"
    best_mime = "image/jpeg"

    def _update(data: bytes, ext: str, mime: str) -> None:
        nonlocal best_data, best_size, best_ext, best_mime
        size = len(data)
        if size <= target_bytes:
            if best_size > target_bytes or size > best_size:
                best_data, best_size, best_ext, best_mime = data, size, ext, mime
        else:
            if best_size > target_bytes and size < best_size:
                best_data, best_size, best_ext, best_mime = data, size, ext, mime

    if has_alpha:
        # 1. Try lossless PNG first
        png_data = _render("PNG", 0, 1.0)
        _update(png_data, ".png", "image/png")
        if best_size <= target_bytes:
            return best_data, best_size, True, best_ext, best_mime

        # 2. Use WebP for lossy compression with alpha preservation
        lo_q, hi_q = 10, 90
        for _ in range(8):
            if lo_q > hi_q:
                break
            q = (lo_q + hi_q) // 2
            c = _render("WEBP", q, 1.0)
            _update(c, ".webp", "image/webp")
            if len(c) <= target_bytes:
                lo_q = q + 1
            else:
                hi_q = q - 1

        if best_size <= target_bytes:
            return best_data, best_size, True, best_ext, best_mime

        # Binary search scale at quality 40
        lo_s, hi_s = 0.1, 1.0
        for _ in range(8):
            s = (lo_s + hi_s) / 2.0
            c = _render("WEBP", 40, s)
            _update(c, ".webp", "image/webp")
            if len(c) <= target_bytes:
                lo_s = s
            else:
                hi_s = s

        return best_data, best_size, best_size <= target_bytes, best_ext, best_mime

    else:
        # Opaque image -> JPEG
        init_data = _render("JPEG", 85, 1.0)
        _update(init_data, ".jpg", "image/jpeg")
        if best_size <= target_bytes:
            return best_data, best_size, True, best_ext, best_mime

        lo_q, hi_q = 5, 85
        for _ in range(12):
            if lo_q > hi_q:
                break
            q = (lo_q + hi_q) // 2
            c = _render("JPEG", q, 1.0)
            _update(c, ".jpg", "image/jpeg")
            if len(c) <= target_bytes:
                lo_q = q + 1
            else:
                hi_q = q - 1

        if best_size <= target_bytes:
            return best_data, best_size, True, best_ext, best_mime

        lo_s, hi_s = 0.1, 1.0
        for _ in range(8):
            s = (lo_s + hi_s) / 2.0
            c = _render("JPEG", 30, s)
            _update(c, ".jpg", "image/jpeg")
            if len(c) <= target_bytes:
                lo_s = s
            else:
                hi_s = s

        return best_data, best_size, best_size <= target_bytes, best_ext, best_mime


@bp.route("/compress-image")
def compress_image_page():
    return render_template(
        "tools/compress_image_plus.html",
        title="Compress Images",
        description="Upload images and set your exact desired output file size.",
        endpoint="/acca-plus/compress-image",
        accept=".jpg,.jpeg,.png,.webp,.bmp",
    )


@bp.route("/compress-image", methods=["POST"])
def compress_image_target():
    files = request.files.getlist("files")
    if not files or not files[0].filename:
        return jsonify(error="Please upload at least one image."), 400

    try:
        target_value = float(request.form.get("target_size", "0") or "0")
        target_unit = (request.form.get("target_unit", "KB") or "KB").upper()
        target_bytes = int(target_value * (1024 * 1024 if target_unit == "MB" else 1024))
    except (ValueError, TypeError):
        return jsonify(error="Invalid target size value."), 400

    if target_bytes <= 0:
        return jsonify(error="Target size must be greater than 0."), 400

    # If single file, return it directly
    if len(files) == 1:
        upload = files[0]
        original_name = upload.filename or "image.jpg"
        stem = os.path.splitext(original_name)[0]
        data = upload.read()
        
        try:
            res_data, achv, met, ext, mime = _compress_image_to_target(data, target_bytes)
        except Exception as e:
            log_error(e, f"compress_image.{original_name}")
            return jsonify(error=f"Failed to compress image: {e}"), 400
            
        out_buf = io.BytesIO(res_data)
        out_buf.seek(0)
        resp = send_file(
            out_buf,
            mimetype=mime,
            as_attachment=True,
            download_name=f"{stem}_compressed{ext}",
        )
        resp.headers["X-Original-Size"] = str(len(data))
        resp.headers["X-Achieved-Size"] = str(achv)
        resp.headers["X-Target-Size"] = str(target_bytes)
        resp.headers["X-Target-Met"] = "1" if met else "0"
        resp.headers["Access-Control-Expose-Headers"] = "X-Original-Size, X-Achieved-Size, X-Target-Size, X-Target-Met"
        return resp

    # If multiple files, return a ZIP
    zip_buf = io.BytesIO()
    log_lines: list[str] = []
    used_names: set[str] = set()
    processed = 0

    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for upload in files:
            orig_name = upload.filename or "unknown.jpg"
            stem = os.path.splitext(orig_name)[0]
            data = upload.read()
            
            try:
                res_data, achv, met, ext, mime = _compress_image_to_target(data, target_bytes)
                new_name = _unique_name(f"{stem}_compressed{ext}", used_names)
                zf.writestr(new_name, res_data)
                
                status = "OK" if met else f"Target not met (achieved {achv} bytes)"
                log_lines.append(f"{orig_name} → {new_name} [{status}]")
                processed += 1
            except Exception as e:
                log_error(e, f"compress_image.{orig_name}")
                fallback_name = _unique_name(orig_name, used_names)
                zf.writestr(fallback_name, data)
                log_lines.append(f"{orig_name} → ERROR: {e} (kept original)")
                
        zf.writestr("_compression_log.txt", "\n".join(log_lines).encode("utf-8"))

    if processed == 0:
        return jsonify(error="No images could be processed."), 400

    zip_buf.seek(0)
    resp = send_file(
        zip_buf,
        mimetype="application/zip",
        as_attachment=True,
        download_name="compressed_images.zip",
    )
    # Give a generic header so the frontend knows it's a zip response
    resp.headers["X-Is-Zip"] = "1"
    resp.headers["Access-Control-Expose-Headers"] = "X-Is-Zip"
    return resp

